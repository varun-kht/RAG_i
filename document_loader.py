"""
Document Loader and Text Chunker module for CILI RAG System.
Supports PDF page tracking, text, markdown, csv, json, docx parsing.
Compatible with Python 3.11+.
"""

import io
import json
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

@dataclass
class Document:
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    doc_id: Optional[str] = None

@dataclass
class TextChunk:
    chunk_id: str
    doc_id: str
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    chunk_index: int = 0


class DocumentLoader:
    """Loads text documents from various file formats and memory streams."""

    @staticmethod
    def load_pdf_stream(file_bytes: bytes, filename: str = "uploaded_document.pdf") -> Document:
        """Extract text from PDF raw bytes stream with page tracking."""
        pages_text = []
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page_num, page in enumerate(reader.pages, 1):
                t = page.extract_text() or ""
                if t.strip():
                    pages_text.append(f"[Page {page_num}]\n{t.strip()}")
        except Exception:
            # Fallback text extraction from raw bytes
            raw = file_bytes.decode("latin1", errors="ignore")
            text_parts = re.findall(r"\((.*?)\)", raw)
            extracted = "\n".join([p for p in text_parts if len(p) > 10])
            pages_text.append(extracted if extracted else "No readable text extracted from PDF.")

        full_content = "\n\n".join(pages_text)
        return Document(
            content=full_content,
            metadata={
                "source": filename,
                "file_type": ".pdf",
                "file_size": len(file_bytes)
            },
            doc_id=filename
        )

    @staticmethod
    def load_file(file_path: str | Path) -> Document:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        metadata = {
            "source": path.name,
            "file_path": str(path.resolve()),
            "file_type": ext,
            "file_size": path.stat().st_size
        }

        if ext == ".pdf":
            with open(path, "rb") as f:
                return DocumentLoader.load_pdf_stream(f.read(), filename=path.name)
        elif ext in [".txt", ".md", ".log", ".py", ".html", ".js", ".json"]:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
        elif ext == ".csv":
            import csv
            lines = []
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(" | ".join(row))
            content = "\n".join(lines)
        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(path)
                content = "\n".join([para.text for para in doc.paragraphs])
            except ImportError:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
        else:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

        return Document(content=content, metadata=metadata, doc_id=path.name)

    @staticmethod
    def load_directory(dir_path: str | Path) -> List[Document]:
        directory = Path(dir_path)
        documents = []
        if not directory.exists() or not directory.is_dir():
            return documents

        supported_extensions = {".txt", ".md", ".pdf", ".csv", ".json", ".docx", ".log", ".py"}
        for file in directory.rglob("*"):
            if file.is_file() and file.suffix.lower() in supported_extensions:
                try:
                    documents.append(DocumentLoader.load_file(file))
                except Exception as e:
                    print(f"[Warning] Failed to load {file}: {e}")
        return documents


class RecursiveTextSplitter:
    """Recursive Text Splitter for breaking text into overlapping chunks."""

    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 100):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.separators = ["\n\n[Page ", "\n\n", "\n", ". ", "? ", "! ", "; ", ", ", " "]

    def split_document(self, document: Document) -> List[TextChunk]:
        raw_chunks = self._split_text(document.content)
        chunks = []
        for i, text in enumerate(raw_chunks):
            chunk_metadata = document.metadata.copy()
            chunk_metadata["chunk_index"] = i
            chunk_id = f"{document.doc_id or 'doc'}_chunk_{i}"
            chunks.append(TextChunk(
                chunk_id=chunk_id,
                doc_id=document.doc_id or "doc",
                text=text.strip(),
                metadata=chunk_metadata,
                chunk_index=i
            ))
        return chunks

    def _split_text(self, text: str) -> List[str]:
        if not text or len(text.strip()) == 0:
            return []

        if len(text) <= self.chunk_size:
            return [text]

        chosen_sep = ""
        for sep in self.separators:
            if sep in text:
                chosen_sep = sep
                break

        if not chosen_sep:
            return [text[i:i + self.chunk_size] for i in range(0, len(text), self.chunk_size - self.chunk_overlap)]

        splits = text.split(chosen_sep)
        chunks = []
        current_chunk = []
        current_length = 0

        for split in splits:
            item = split + (chosen_sep if not chosen_sep.startswith("\n\n[Page ") else "")
            if current_length + len(item) > self.chunk_size and current_chunk:
                joined = "".join(current_chunk).strip()
                if joined:
                    chunks.append(joined)
                
                overlap_acc = []
                overlap_len = 0
                for prev in reversed(current_chunk):
                    if overlap_len + len(prev) <= self.chunk_overlap:
                        overlap_acc.insert(0, prev)
                        overlap_len += len(prev)
                    else:
                        break
                current_chunk = overlap_acc
                current_length = overlap_len

            current_chunk.append(item)
            current_length += len(item)

        if current_chunk:
            final_text = "".join(current_chunk).strip()
            if final_text:
                chunks.append(final_text)

        return chunks
